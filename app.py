import io
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.express as px
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
import streamlit as st

# ================= ========================================
# 1. إعدادات الصفحة وواجهة المستخدم (Streamlit RTL)
# ==========================================================
st.set_page_config(
    page_title="نظام إدارة وصيانة المولدات الذكي",
    page_icon="⚡",
    layout="wide",
)

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700&display=swap');
    html, body, [class*="css"]  {
        font-family: 'Cairo', sans-serif;
        direction: rtl;
        text-align: right;
    }
    .stMetric {
        background-color: #f8f9fa;
        padding: 15px;
        border-radius: 10px;
        border-right: 5px solid #1f4e79;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("⚡ نظام متابعة صيانة المولدات والجاهزية التشغيلية")
st.caption(
    "Addoma Trading Services — تتبع ساعات تشغيل الزيت والفلاتر، تقييم جاهزية الورديات، والتصدير الاحترافي"
)
st.markdown("---")

# ==========================================================
# 2. إدارة البيانات والتنسيق (Session State)
# ==========================================================
if "generators" not in st.session_state:
    st.session_state.generators = [
        {
            "id": "GEN-001",
            "location": "موقع التعدين - الموقع A",
            "current_hrs": 4250,
            "last_change_hrs": 4000,
            "max_oil_hrs": 250,
            "next_shift_hrs": 200,
        },
        {
            "id": "GEN-002",
            "location": "موقع التعدين - الموقع B",
            "current_hrs": 3120,
            "last_change_hrs": 3000,
            "max_oil_hrs": 250,
            "next_shift_hrs": 100,
        },
        {
            "id": "GEN-003",
            "location": "جامعة البيان - المولد الرئيسي",
            "current_hrs": 1850,
            "last_change_hrs": 1600,
            "max_oil_hrs": 250,
            "next_shift_hrs": 50,
        },
        {
            "id": "GEN-004",
            "location": "غرفة التبريد - المستشفى",
            "current_hrs": 5400,
            "last_change_hrs": 5180,
            "max_oil_hrs": 250,
            "next_shift_hrs": 40,
        },
        {
            "id": "GEN-005",
            "location": "الموقع الزراعي - المحورية",
            "current_hrs": 2900,
            "last_change_hrs": 2660,
            "max_oil_hrs": 250,
            "next_shift_hrs": 60,
        },
    ]

# ==========================================================
# 3. القائمة الجانبية (إضافة / تحديث المولدات)
# ==========================================================
with st.sidebar:
    st.header("➕ إضافة / تحديث بيانات مولد")
    with st.form("gen_form", clear_on_submit=True):
        gen_id = st.text_input("معرف المولد (ID)", value="GEN-006")
        location = st.text_input("الموقع / القسم", value="الموقع الصناعي")
        current_hrs = st.number_input(
            "ساعات التشغيل الحالية", min_value=0, value=3500
        )
        last_change_hrs = st.number_input(
            "ساعات آخر غيار زيت", min_value=0, value=3300
        )
        max_oil_hrs = st.number_input(
            "العمر الافتراضي للزيت (ساعة)", min_value=1, value=250
        )
        next_shift_hrs = st.number_input(
            "ساعات وردية التشغيل القادمة", min_value=0, value=120
        )

        submitted = st.form_submit_button("حفظ وتحديث السجل")
        if submitted:
            existing_idx = next(
                (
                    i
                    for i, g in enumerate(st.session_state.generators)
                    if g["id"] == gen_id
                ),
                None,
            )
            new_data = {
                "id": gen_id,
                "location": location,
                "current_hrs": current_hrs,
                "last_change_hrs": last_change_hrs,
                "max_oil_hrs": max_oil_hrs,
                "next_shift_hrs": next_shift_hrs,
            }
            if existing_idx is not None:
                st.session_state.generators[existing_idx] = new_data
                st.success(f"تم تحديث بيانات المولد {gen_id} بنجاح!")
            else:
                st.session_state.generators.append(new_data)
                st.success(f"تمت إضافة المولد {gen_id} بنجاح!")

# ==========================================================
# 4. المعالجة الحسابية والتحليل
# ==========================================================
df = pd.DataFrame(st.session_state.generators)

if not df.empty:
    df["used_hrs"] = df["current_hrs"] - df["last_change_hrs"]
    df["remaining_hrs"] = df["max_oil_hrs"] - df["used_hrs"]

    df["shift_readiness"] = df.apply(
        lambda row: (
            "جاهز للوردية"
            if row["remaining_hrs"] >= row["next_shift_hrs"]
            else "غير جاهز - يتطلب صيانة"
        ),
        axis=1,
    )


    def get_status(rem):
        if rem <= 0:
            return "تجاوز الخدمة!"
        elif rem <= 50:
            return "صيانة قريبة"
        else:
            return "ممتاز"


    df["main_status"] = df["remaining_hrs"].apply(get_status)

    # ==========================================================
    # 5. عرض البطاقات الإحصائية (KPIs)
    # ==========================================================
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("إجمالي المولدات", len(df))
    col2.metric(
        "جاهزة للورديات القادمة",
        len(df[df["shift_readiness"] == "جاهز للوردية"]),
    )
    col3.metric(
        "تتطلب صيانة عاجلة", len(df[df["main_status"] == "تجاوز الخدمة!"])
    )
    col4.metric("اقتراب موعد الصيانة", len(df[df["main_status"] == "صيانة قريبة"]))

    st.markdown("---")

    # ==========================================================
    # 6. الرسوم البيانية التوضيحية (Plotly)
    # ==========================================================
    st.subheader("📊 الرسوم البيانية والتحليل التفاعلي")
    chart_col1, chart_col2 = st.columns(2)

    with chart_col1:
        fig_bar = px.bar(
            df,
            x="id",
            y=["used_hrs", "remaining_hrs"],
            title="مقارنة الساعات المستهلكة مقابل الساعات المتبقية للخدمة",
            labels={
                "value": "الساعات",
                "id": "معرف المولد",
                "variable": "المؤشر",
            },
            color_discrete_sequence=["#d9534f", "#28a745"],
            barmode="group",
        )
        st.plotly_chart(fig_bar, use_container_width=True)

    with chart_col2:
        fig_pie = px.pie(
            df,
            names="shift_readiness",
            title="توزيع جاهزية المولدات للورديات القادمة",
            color="shift_readiness",
            color_discrete_map={
                "جاهز للوردية": "#28a745",
                "غير جاهز - يتطلب صيانة": "#dc3545",
            },
        )
        st.plotly_chart(fig_pie, use_container_width=True)

    # ==========================================================
    # 7. عرض الجدول التفصيلي
    # ==========================================================
    st.subheader("📋 سجل متابعة المولدات وصيانة الزيت")
    display_df = df.rename(
        columns={
            "id": "معرف المولد",
            "location": "الموقع",
            "current_hrs": "ساعات التشغيل الحالية",
            "last_change_hrs": "آخر غيار زيت",
            "max_oil_hrs": "عمر الزيت",
            "used_hrs": "الساعات المستهلكة",
            "remaining_hrs": "الساعات المتبقية",
            "next_shift_hrs": "الوردية القادمة",
            "shift_readiness": "جاهزية الوردية",
            "main_status": "حالة الصيانة",
        }
    )
    st.dataframe(display_df, use_container_width=True)

    # ==========================================================
    # 8. خيارات تصدير التقارير (Excel مطور & PDF)
    # ==========================================================
    st.markdown("---")
    st.subheader("📥 تصدير التقارير والبيانات")

    exp_col1, exp_col2 = st.columns(2)

    # 🟢 أ. دالة توليد ملف إكسل المطور (معادلات + رسوم بيانية + ورقتا عمل)
    def generate_advanced_excel():
        wb = openpyxl.Workbook()

        # Sheet 1: Dashboard
        ws_dash = wb.active
        ws_dash.title = "لوحة التحكم والتقرير"
        ws_dash.views.sheetView[0].rightToLeft = True

        # Sheet 2: Data
        ws_data = wb.create_sheet(title="سجل المولدات والصيانة")
        ws_data.views.sheetView[0].rightToLeft = True

        HEADER_FILL = PatternFill(
            start_color="1F4E79", end_color="1F4E79", fill_type="solid"
        )
        HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E79")
        SECTION_FONT = Font(name="Calibri", size=13, bold=True, color="1F4E79")
        REGULAR_FONT = Font(name="Calibri", size=11)
        BORDER = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # 1. إدخال ورقة البيانات التفصيلية
        headers_data = [
            "معرف المولد",
            "موقع المولد",
            "ساعات التشغيل الحالية",
            "قراءة آخر غيار زيت",
            "العمر الافتراضي للزيت",
            "الساعات المستهلكة",
            "الساعات المتبقية للخدمة",
            "ساعات وردية التشغيل القادمة",
            "حالة الجاهزية للوردية",
            "حالة الصيانة الحالية",
        ]
        ws_data.append(headers_data)
        for col_num, h in enumerate(headers_data, 1):
            c = ws_data.cell(row=1, column=col_num)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

        for idx, row in df.iterrows():
            r = idx + 2
            row_data = [
                row["id"],
                row["location"],
                row["current_hrs"],
                row["last_change_hrs"],
                row["max_oil_hrs"],
                f"=C{r}-D{r}",
                f"=E{r}-F{r}",
                row["next_shift_hrs"],
                f'=IF(G{r}>=H{r}, "جاهز للوردية", "غير جاهز - يتطلب صيانة")',
                f'=IF(G{r}<=0, "تجاوز الخدمة!", IF(G{r}<=50, "صيانة قريبة", "ممتاز"))',
            ]
            ws_data.append(row_data)

        for row in ws_data.iter_rows(
            min_row=2, max_row=len(df) + 1, min_col=1, max_col=10
        ):
            for cell in row:
                cell.font = REGULAR_FONT
                cell.border = BORDER
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

        # 2. إدخال ورقة لوحة التحكم (Dashboard)
        ws_dash.merge_cells("A1:F1")
        ws_dash["A1"] = (
            "تقرير أداء وصيانة المولدات — Addoma Trading Services"
        )
        ws_dash["A1"].font = TITLE_FONT
        ws_dash["A1"].alignment = Alignment(
            horizontal="center", vertical="center"
        )

        ws_dash.merge_cells("A2:F2")
        ws_dash["A2"] = (
            "نظام التتبع الذكي لساعات التشغيل، صيانة الزيت، وجاهزية الورديات"
        )
        ws_dash["A2"].font = Font(
            name="Calibri", size=11, italic=True, color="595959"
        )
        ws_dash["A2"].alignment = Alignment(
            horizontal="center", vertical="center"
        )

        ws_dash["A4"] = "ملخص حالة المولدات والساعات المتبقية"
        ws_dash["A4"].font = SECTION_FONT

        dash_headers = [
            "معرف المولد",
            "الموقع",
            "ساعات التشغيل",
            "الساعات المستهلكة",
            "الساعات المتبقية",
            "حالة الخدمة",
        ]
        for col_idx, h in enumerate(dash_headers, 1):
            cell = ws_dash.cell(row=5, column=col_idx)
            cell.value = h
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, row in df.iterrows():
            r = idx + 2
            dash_r = idx + 6
            ws_dash.cell(row=dash_r, column=1, value=f"='سجل المولدات والصيانة'!A{r}")
            ws_dash.cell(row=dash_r, column=2, value=f"='سجل المولدات والصيانة'!B{r}")
            ws_dash.cell(row=dash_r, column=3, value=f"='سجل المولدات والصيانة'!C{r}")
            ws_dash.cell(row=dash_r, column=4, value=f"='سجل المولدات والصيانة'!F{r}")
            ws_dash.cell(row=dash_r, column=5, value=f"='سجل المولدات والصيانة'!G{r}")
            ws_dash.cell(row=dash_r, column=6, value=f"='سجل المولدات والصيانة'!J{r}")

            for c in range(1, 7):
                cell = ws_dash.cell(row=dash_r, column=c)
                cell.font = REGULAR_FONT
                cell.border = BORDER
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

        # إضافة الرسم البياني داخل ملف الإكسل
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "مقارنة الساعات المستهلكة مقابل الساعات المتبقية"
        chart.y_axis.title = "الساعات"
        chart.x_axis.title = "المولدات"

        data_ref = Reference(
            ws_dash,
            min_col=4,
            min_row=5,
            max_col=5,
            max_row=len(df) + 5,
        )
        cats_ref = Reference(
            ws_dash, min_col=1, min_row=6, max_row=len(df) + 5
        )
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 16
        chart.height = 10

        ws_dash.add_chart(chart, f"A{len(df) + 8}")

        # ضبط عرض الأعمدة تلقائياً
        for ws in [ws_dash, ws_data]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    with exp_col1:
        excel_bytes = generate_advanced_excel()
        st.download_button(
            label="📊 تصدير البيانات إلى Excel (تقرير متطور مع رسوم ومعادلات)",
            data=excel_bytes,
            file_name="تقرير_صيانة_المولدات_المطور.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # 🔴 ب. دالة توليد تقرير PDF المطور
    with exp_col2:

        def generate_pdf():
            pdf_buffer = io.BytesIO()
            doc = SimpleDocTemplate(
                pdf_buffer,
                pagesize=landscape(A4),
                rightMargin=20,
                leftMargin=20,
                topMargin=20,
                bottomMargin=20,
            )
            elements = []
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "TitleStyle",
                parent=styles["Heading1"],
                alignment=1,
                fontSize=18,
                textColor=colors.HexColor("#1F4E79"),
            )

            elements.append(
                Paragraph(
                    "<b>Addoma Trading Services — Generator Maintenance Report</b>",
                    title_style,
                )
            )
            elements.append(Spacer(1, 10))
            elements.append(
                HRFlowable(
                    width="100%", thickness=2, color=colors.HexColor("#1F4E79")
                )
            )
            elements.append(Spacer(1, 15))

            table_data = [[col for col in display_df.columns]]
            for _, row in display_df.iterrows():
                table_data.append([str(val) for val in row.values])

            pdf_table = Table(table_data, repeatRows=1)
            pdf_table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#1F4E79"),
                        ),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                        (
                            "GRID",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.HexColor("#D9D9D9"),
                        ),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F2F2F2")],
                        ),
                    ]
                )
            )

            elements.append(pdf_table)
            doc.build(elements)
            return pdf_buffer.getvalue()


        pdf_bytes = generate_pdf()
        st.download_button(
            label="📄 تصدير التقرير الرسمي إلى PDF",
            data=pdf_bytes,
            file_name="تقرير_صيانة_المولدات.pdf",
            mime="application/pdf",
        )
else:
    st.info("الرجاء إضافة مولدات من القائمة الجانبية لعرض التقرير والرسوم البيانية.")
